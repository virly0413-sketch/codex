"""Read-only schedule/share audit. Never uploads, moves or deletes videos."""
import datetime as dt
import difflib
import hashlib
import os
from pathlib import Path
import tempfile

from baidu_download import read_csv, date_columns, row_language, LANGUAGES, file_language, sha256_file
from network_upload import inventory, matches, title_key
from prepare import normalize


def schedule_tasks(path, selections, start=None, days=None, infer_blank=False):
    table=read_csv(path)
    if not table:raise ValueError('排片表为空')
    dates=date_columns(table,infer_blank)
    end=None
    if start is not None:
        if not start:raise ValueError('请填写开始日期，或明确选择全部日期')
        day=dt.date.fromisoformat(start)
        if not isinstance(days,int) or not 1<=days<=366:raise ValueError('天数须为1—366')
        end=(day+dt.timedelta(days=days-1)).isoformat()
        if start not in {d[0] for d in dates}:raise ValueError('开始日期不在日期表头中，请核对原表；不会擅自顺延空白日期。')
    if not selections:raise ValueError('请选择具体排片行')
    tasks=[];seen=set()
    from openpyxl.utils import get_column_letter
    for selection in selections:
        number=selection['row'];language=selection['language']
        if number in seen:raise ValueError('同一行不能重复选择')
        seen.add(number)
        if not isinstance(number,int) or not 2<=number<=len(table):raise ValueError('排片行不存在，请重新读取')
        row=table[number-1]
        if language not in LANGUAGES or row_language(row[0])!=language:raise ValueError('行语种不明确，请先在CSV的行名标明正确语种')
        if any(c.strip() for c in row[len(table[0]):]):raise ValueError('剧名超出日期表头，请补全原表日期')
        for date,inferred,column in dates:
            if not date or (start and date<start) or (end and date>end):continue
            cell=row[column-1] if column<=len(row) else ''
            for slot,title in enumerate(cell.split('/'),1):
                title=title.strip()
                if not title:continue
                tasks.append(dict(date=date,language=language,title=title,candidates=[],source_row_label=row[0],
                                  source_cell=get_column_letter(column)+str(number),slot=slot,inferred_date=inferred))
    if not tasks:raise ValueError('所选范围没有可核对排片；请检查行、日期和空白日期设置')
    return tasks


def audit(tasks, files, stop=None):
    result=[]
    for task in tasks:
        if stop and stop.is_set():raise InterruptedError('已停止核对，未导出未完成名单')
        found=matches(task,files);same=[];possible=[];key=normalize(task['title'])
        if not found:
            for f in files:
                labels=[title_key(label) for label in [f['name']]+f.get('parents',[])]
                if key in labels:same.append(f)
                elif any(len(label)>=4 and (key in label or label in key or difflib.SequenceMatcher(None,key,label).ratio()>=0.78) for label in labels):
                    possible.append(f)
        status=('已找到' if len(found)==1 else '多个对应视频' if found else
                '同剧其他语种或语种未标明' if same else '疑似同剧，需核对命名' if possible else '未匹配对应视频')
        result.append(dict(task,status=status,files=found,related=same or possible[:10]))
    return result


def export_report(path, results, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book=Workbook()
    headers=['排片日期','频道/原表行名','语种','剧名','原表位置','核对结果','对应视频数量','匹配路径','其他语种/疑似路径','提示']
    def add(name,rows):
        ws=book.create_sheet(name);ws.append(headers)
        for item in rows:
            path_text='\n'.join(f['path'] for f in item['files'])
            related='\n'.join((file_language(f) or '语种不明')+' | '+f['path'] for f in item['related'])
            note='多个文件可能是分集或不同版本，请人工选定' if len(item['files'])>1 else ('仅按名称/目录及语言标记匹配，未检验音轨或完整性' if item['files'] else '未匹配不等于网盘中绝对没有；请核对命名、语种和搜索根目录')
            if item['inferred_date']:note+='；日期为用户勾选顺延后推算'
            ws.append([item['date'],item['source_row_label'],item['language'],item['title'],item['source_cell'],item['status'],str(len(item['files'])),path_text[:32000],related[:32000],note])
            for c in ws[ws.max_row]:c.data_type='s';c.alignment=Alignment(wrap_text=True,vertical='top')
            if item['files']:ws.cell(ws.max_row,8).hyperlink=item['files'][0]['path']
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
        for col,width in {'A':15,'B':26,'C':14,'D':45,'E':14,'F':34,'G':18,'H':90,'I':100,'J':80}.items():ws.column_dimensions[col].width=width
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    book.remove(book.active)
    add('未匹配对应语种',[r for r in results if not r['files']])
    add('多个对应视频',[r for r in results if len(r['files'])>1])
    add('已找到',[r for r in results if len(r['files'])==1])
    add('全部排片',results)
    ws=book.create_sheet('核对范围');ws.append(['项目','内容'])
    for key,value in metadata.items():
        ws.append([key,str(value)]);ws.cell(ws.max_row,2).data_type='s'
    ws.append(['结果含义','未匹配表示本次所选目录中没有满足剧名和语言规则的非空视频，不是对所有存储位置或实际音轨的不存在证明。'])
    ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=130
    target=Path(path);target.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(dir=target.parent,suffix='.xlsx');os.close(fd)
    try:book.save(tmp);os.replace(tmp,target)
    finally:
        book.close();Path(tmp).unlink(missing_ok=True)


def check(schedule, share, selections, output, start=None, days=None, infer_blank=False, stop=None, progress=print):
    before=sha256_file(schedule)
    tasks=schedule_tasks(schedule,selections,start,days,infer_blank)
    progress('正在递归读取共享目录…')
    files=inventory(share,stop)
    progress(f'已读取{len(files)}个非空视频，正在按剧名和语种核对…')
    results=audit(tasks,files,stop)
    if before!=sha256_file(schedule):raise ValueError('核对期间排片表发生变化，请重新扫描')
    if stop and stop.is_set():raise InterruptedError('已停止，未导出未完成名单')
    export_report(output,results,{'扫描北京时间':dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(timespec='seconds'),
        '排片表':str(Path(schedule).resolve()),'排片SHA256':before,'共享根目录':str(Path(share).resolve()),
        '选择行':'; '.join(f"{r['row']}:{r['language']}" for r in selections),'开始日期':start or '所有明确日期',
        '天数':days if start else '全部','允许空白日期顺延':infer_blank,'非空视频文件数':len(files),
        '排片任务数':len(results),'未匹配对应语种':sum(not r['files'] for r in results),
        '多个对应视频':sum(len(r['files'])>1 for r in results),
        '完整性':'所选目录递归读取已完成；不可访问子目录会停止，不生成缺片结论。符号链接不遍历。'})
    return results
