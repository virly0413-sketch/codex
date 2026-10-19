import csv
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import MagicMock

from openpyxl import load_workbook

from baidu_browser import BaiduBrowser, DownloadUnavailable, listing, private_error, validate_video
from baidu_download import (build_tasks, destination, file_language, select_files, share_identity,
                            share_link, suffix_language)
from baidu_runner import run_downloads


class DownloadTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.task=dict(id='task',date='2026-10-20',language='俄语',title='渔夫从神奇龙珠开始起飞',
            code='CODE',share=dict(url='https://pan.baidu.com/s/demo',password='abcd'),candidates=[])
        self.payload=b'\x00\x00\x00\x18ftypisom'+b'example video bytes'*30
        self.item=dict(id='1001',name='一介渔夫靠龙珠号令沧海俄语.mp4',path='/files/俄语.mp4',
                       isdir=False,size=len(self.payload),parents=[])
        self.browser=MagicMock()
        self.browser.inventory.return_value=[self.item]
        def download(share,item,path):
            if Path(path).exists():raise ValueError('既有文件未校验，不覆盖')
            Path(path).parent.mkdir(parents=True,exist_ok=True)
            Path(path).write_bytes(self.payload)
        self.browser.download.side_effect=download

    def tearDown(self):self.temp.cleanup()

    def execute(self,tasks=None,scan_only=False):
        return run_downloads(tasks or [self.task],{'俄语':self.root/'俄语','英语':self.root/'英语'},
            self.root/'work',self.browser,scan_only=scan_only,log=lambda _:None)

    def test_suffix_filenames_all_languages_and_renamed_drama(self):
        for language in ['俄语','英语','西语','菲律宾语','韩语']:
            self.assertEqual(suffix_language('另一个发行名'+language+'.MP4'),language)
        self.assertEqual(suffix_language('另一个发行名【俄语】_1080P.mp4'),'俄语')
        self.assertIsNone(suffix_language('这个剧名提到俄语但是中文.mp4'))
        self.assertIsNone(suffix_language('全集英语俄语.mp4'))
        self.assertIsNone(suffix_language('全集非俄语.mp4'))

    def test_folder_language_and_explicit_file_priority(self):
        item=dict(self.item,name='第01集.mp4',parents=['发行名俄语','正片'])
        self.assertEqual(file_language(item),'俄语')
        item['name']='正片英语.mp4';self.assertEqual(file_language(item),'英语')
        item['name']='全集英语俄语.mp4';self.assertIsNone(file_language(item))
        item['name']='01.mp4';item['parents']=['发行名俄语','英语俄语'];self.assertIsNone(file_language(item))

    def test_select_only_video_and_requested_language(self):
        items=[self.item,dict(self.item,id='2',name='正片英语.mp4'),
            dict(self.item,id='3',name='字幕俄语.srt'),dict(self.item,id='4',name='俄语',isdir=True),
            dict(self.item,id='5',name='未知语言.mp4')]
        self.assertEqual(select_files(items,'俄语'),[self.item])

    def test_share_parser_and_init_identity(self):
        s=share_link('链接：https://pan.baidu.com/s/demo?pwd=Ab12 提取码：Ab12')
        self.assertEqual(s['password'],'Ab12')
        self.assertIsNone(share_link('https://example.com/s/demo'))
        self.assertIsNone(share_link('https://pan.baidu.com/s/a https://pan.baidu.com/s/b'))
        self.assertNotEqual(share_identity({'url':'https://pan.baidu.com/share/init?surl=a'}),
            share_identity({'url':'https://pan.baidu.com/share/init?surl=b'}))
        self.assertIsNone(share_link('https://pan.baidu.com/share/init?pwd=abcd'))

    def test_plan_infers_dates_and_keeps_version_ambiguous(self):
        schedule=self.root/'schedule.csv';resources=self.root/'resources.csv'
        with schedule.open('w',encoding='utf-8-sig',newline='') as f:
            csv.writer(f).writerows([['日期','2026.10.25','',''],['俄语男频','甲','乙','甲']])
        rows=[['资源表'],['剧名','代码']]
        for title,code in [('《甲》','A'),('《乙》','B1'),('《乙》','B2')]:
            row=['']*30;row[0]=title;row[1]=code;row[29]='https://pan.baidu.com/s/demo?pwd=abcd';rows.append(row)
        with resources.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
        tasks=build_tasks(schedule,resources,['俄语'],start='2026-10-26')
        self.assertEqual([x['date'] for x in tasks],['2026-10-26','2026-10-27'])
        self.assertEqual(tasks[0]['status'],'待选择版本代码')
        self.assertIsNone(tasks[0]['share'])
        self.assertTrue(all(x['inferred_date'] for x in tasks))

    def test_real_file_writes_separate_languages_and_cached_listing(self):
        en=dict(self.task,id='english',language='英语')
        english_item=dict(self.item,id='2002',name='一介渔夫靠龙珠号令沧海英语.mp4')
        self.browser.inventory.return_value=[self.item,english_item]
        records=self.execute([self.task,en])
        self.assertEqual([r['状态'] for r in records],['下载完成','下载完成'])
        self.assertTrue(Path(records[0]['本地文件']).is_relative_to(self.root/'俄语'))
        self.assertTrue(Path(records[1]['本地文件']).is_relative_to(self.root/'英语'))
        self.browser.inventory.assert_called_once()
        self.assertEqual(self.browser.download.call_count,2)

    def test_complete_file_skipped_only_after_hash_check(self):
        first=self.execute()[0]
        self.assertEqual(self.execute()[0]['状态'],'已完成，校验后跳过')
        self.assertEqual(self.browser.download.call_count,1)
        path=Path(first['本地文件']);path.write_bytes(b'x'*len(self.payload))
        self.assertEqual(self.execute()[0]['状态'],'下载失败')
        self.assertEqual(self.browser.download.call_count,2)

    def test_scan_only_never_downloads_or_makes_output_directory(self):
        self.assertEqual(self.execute(scan_only=True)[0]['状态'],'仅扫描匹配')
        self.browser.download.assert_not_called()
        self.assertFalse((self.root/'俄语').exists())

    def test_missing_link_or_code_never_opens_share(self):
        tasks=[dict(self.task,id='missing',share=None),dict(self.task,id='version',code='',candidates=[{}])]
        records=self.execute(tasks)
        self.assertEqual(len(records),2);self.browser.inventory.assert_not_called()

    def test_failed_or_client_required_is_not_success(self):
        self.browser.download.side_effect=ValueError('failure')
        self.assertEqual(self.execute()[0]['状态'],'下载失败')
        self.assertFalse((self.root/'work'/'downloads.json').exists())
        self.browser.download.side_effect=DownloadUnavailable('需要客户端')
        with self.assertRaises(DownloadUnavailable):self.execute()
        book=load_workbook(self.root/'work'/'download_report.xlsx')
        try:self.assertEqual(book.active['F2'].value,'需在网页处理或使用客户端')
        finally:book.close()

    def test_no_matching_language_does_not_fallback(self):
        self.browser.inventory.return_value=[dict(self.item,name='片名英语.mp4')]
        self.assertEqual(self.execute()[0]['状态'],'未找到明确语种视频')
        self.browser.download.assert_not_called()

    def test_short_file_and_html_zip_rejected(self):
        path=self.root/'test.mp4';path.write_bytes(b'<html>login</html>')
        with self.assertRaises(ValueError):validate_video(path,'test.mp4',0)
        path.write_bytes(self.payload)
        with self.assertRaises(ValueError):validate_video(path,'test.zip',len(self.payload))
        with self.assertRaises(ValueError):validate_video(path,'test.mp4',999999)
        validate_video(path,'test.mp4',len(self.payload))

    def test_destination_stays_in_flat_language_folder_and_cannot_traverse(self):
        target=destination(dict(self.task,title='../../恶意:标题'),dict(self.item,name='..\\evil俄语.mp4'),self.root/'俄语')
        self.assertEqual(target.parent,self.root/'俄语')
        self.assertNotIn(':',target.name);self.assertNotIn('\\',target.name)
        other=destination(self.task,dict(self.item,id='other'),self.root/'俄语')
        self.assertNotEqual(target,other)

    def test_report_contains_no_share_secrets_or_executable_formula(self):
        self.task['title']='=HYPERLINK("bad")';self.execute(scan_only=True)
        book=load_workbook(self.root/'work'/'download_report.xlsx')
        try:
            self.assertEqual(book.active['C2'].data_type,'s')
            text=str(list(book.active.values))
            self.assertNotIn('pan.baidu.com',text);self.assertNotIn('abcd',text)
        finally:book.close()
        self.assertNotIn('secret',private_error(ValueError('URL https://pan.baidu.com/s/demo?pwd=secret')))

    def test_listing_error_and_repeated_pagination_do_not_become_success(self):
        with self.assertRaises(ValueError):listing({'errno':-9,'list':[]})
        with self.assertRaises(ValueError):listing({'errno':0,'unknown':[]})
        b=BaiduBrowser(self.root/'profile',lambda _:None,lambda:False)
        b.template='https://pan.baidu.com/share/list?uk=1&shareid=2'
        response=MagicMock();response.ok=True
        response.json.return_value={'errno':0,'list':[dict(fs_id=1,server_filename='片名俄语.mp4',isdir=0,size=3)]}
        b.page=MagicMock();b.page.request.get.return_value=response
        with self.assertRaisesRegex(ValueError,'重复页'):b._list_dir('/')

    def test_browser_adapter_delivers_only_valid_video_after_reload(self):
        browser=BaiduBrowser(self.root/'profile',lambda _:None,lambda:False)
        browser.page=MagicMock()
        row=MagicMock();direct=row.locator.return_value.filter.return_value
        direct.count.return_value=1
        browser._file_row=MagicMock(return_value=row)
        download=MagicMock();download.suggested_filename=self.item['name'];download.failure.return_value=None
        download.save_as.side_effect=lambda p:Path(p).write_bytes(self.payload)
        browser.page.expect_download.return_value.__enter__.return_value.value=download
        target=self.root/'result.mp4'
        browser.download(self.task['share'],self.item,target)
        self.assertEqual(target.read_bytes(),self.payload)
        browser.page.reload.assert_called_once()
        direct.click.assert_called_once()
        self.assertFalse(target.with_suffix('.mp4.baidupart').exists())

    def test_browser_adapter_refuses_zip_and_preserves_existing_files(self):
        browser=BaiduBrowser(self.root/'profile',lambda _:None,lambda:False);browser.page=MagicMock()
        row=MagicMock();row.locator.return_value.filter.return_value.count.return_value=1
        browser._file_row=MagicMock(return_value=row)
        download=MagicMock();download.suggested_filename='all_languages.zip'
        browser.page.expect_download.return_value.__enter__.return_value.value=download
        target=self.root/'result.mp4'
        with self.assertRaises(ValueError):browser.download(self.task['share'],self.item,target)
        download.cancel.assert_called_once();download.save_as.assert_not_called()
        self.assertFalse(target.exists())
        target.write_bytes(b'user file')
        with self.assertRaises(ValueError):browser.download(self.task['share'],self.item,target)
        self.assertEqual(target.read_bytes(),b'user file')

    def test_browser_adapter_refuses_multiple_selected_files(self):
        browser=BaiduBrowser(self.root/'profile',lambda _:None,lambda:False);browser.page=MagicMock()
        row=MagicMock();direct=MagicMock();direct.count.return_value=0
        checkbox=MagicMock();checkbox.count.return_value=1
        row.locator.return_value.filter.side_effect=[direct,checkbox]
        browser._file_row=MagicMock(return_value=row)
        browser.page.locator.return_value.inner_text.return_value='已选中2个文件'
        with self.assertRaisesRegex(ValueError,'只选中1个'):browser.download(self.task['share'],self.item,self.root/'bad.mp4')
        browser.page.expect_download.assert_not_called()

    def test_listing_finishes_on_empty_next_page(self):
        browser=BaiduBrowser(self.root/'profile',lambda _:None,lambda:False)
        browser.template='https://pan.baidu.com/share/list?shareid=2&uk=1';browser.page=MagicMock()
        response=MagicMock();response.ok=True
        response.json.side_effect=[{'errno':0,'list':[dict(fs_id=1,server_filename='片名俄语.mp4',isdir=0,size=3)]},
                                  {'errno':0,'list':[]}]
        browser.page.request.get.return_value=response
        self.assertEqual(len(browser._list_dir('/')),1)
        self.assertEqual(browser.page.request.get.call_count,2)


if __name__=='__main__':unittest.main()
