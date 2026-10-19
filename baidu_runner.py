"""Schedule-bound downloads, completed-file verification and local Excel reports."""
import hashlib
import json
from pathlib import Path

from baidu_browser import DownloadUnavailable, private_error
from baidu_download import (atomic_json, destination, ready_status, select_files,
                            sha256_file, share_identity, verified_existing)

HEADERS = ['日期','语种','剧名','剧集代码','网盘文件名','状态','本地文件','说明','排片行','原表单元格']


def write_report(path, records):
    import os
    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(); sheet = book.active; sheet.title = '下载记录'
    sheet.append(HEADERS)
    for row in records:
        sheet.append([row.get(k, '') for k in HEADERS])
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in sheet[1]:
        cell.font = Font(bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor='17365D')
    for column,width in {'A':16,'B':14,'C':40,'D':22,'E':55,'F':28,'G':70,'H':70}.items():
        sheet.column_dimensions[column].width=width
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    fd,temp=tempfile.mkstemp(dir=path.parent,suffix='.xlsx');os.close(fd)
    try:
        book.save(temp);os.replace(temp,path)
    except PermissionError:
        raise PermissionError('请关闭已打开的下载记录Excel，然后重试。') from None
    finally:
        book.close();Path(temp).unlink(missing_ok=True)


def run_downloads(tasks, folders, work, browser, scan_only=False, cancelled=lambda:False, log=print):
    work=Path(work);work.mkdir(parents=True,exist_ok=True)
    state_path=work/'downloads.json'
    state=json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
    records=[];cache={}
    def add(task,status,item=None,path='',note=''):
        records.append(dict(zip(HEADERS,[task['date'],task['language'],task['title'],task.get('code',''),
            item['name'] if item else '',status,str(path),note,task.get('source_row_label',''),task.get('source_cell','')])))
        write_report(work/'download_report.xlsx',records)
        log(f"{task['language']} · {task['title']} · {status}" + ((' · '+item['name']) if item else ''))
    for task in tasks:
        if cancelled():break
        status=ready_status(task)
        if status!='待扫描':
            add(task,status);continue
        if task['language'] not in folders or not str(folders[task['language']]).strip():
            add(task,'未设置该语种目录');continue
        key=(share_identity(task['share']),task['share'].get('password',''))
        try:
            if key not in cache:cache[key]=browser.inventory(task['share'])
            matches=select_files(cache[key],task['language'])
            if not matches:
                add(task,'未找到明确语种视频',note='不回退到其他语种或整包下载。');continue
            for item in matches:
                if cancelled():break
                folder=Path(folders[task['language']]).resolve()
                target=destination(task,item,folder)
                identity=hashlib.sha256((share_identity(task['share'])+'|'+item['id']+'|'+
                    task['language']+'|'+str(folder)).encode()).hexdigest()
                previous=state.get(identity)
                if previous and Path(previous['path']).parent==folder and verified_existing(previous['path'],previous):
                    add(task,'已完成，校验后跳过',item,previous['path']);continue
                if scan_only:
                    add(task,'仅扫描匹配',item,target);continue
                try:
                    log('开始下载：'+task['language']+' · '+task['title']+' · '+item['name'])
                    browser.download(task['share'],item,target)
                    if not target.is_file() or target.stat().st_size<=0:
                        raise ValueError('下载器未交付有效本地文件，未标为成功。')
                    if item['size']>0 and target.stat().st_size!=item['size']:
                        raise ValueError('文件大小与目录记录不一致，未标为成功。')
                    state[identity]=dict(path=str(target),bytes=target.stat().st_size,sha256=sha256_file(target))
                    atomic_json(state_path,state)
                    add(task,'下载完成',item,target)
                except DownloadUnavailable as exc:
                    add(task,'需在网页处理或使用客户端',item,note=private_error(exc))
                    raise
                except InterruptedError:
                    raise
                except Exception as exc:
                    add(task,'下载失败',item,note=private_error(exc))
        except DownloadUnavailable:
            raise
        except InterruptedError as exc:
            add(task,'已停止，未完成',note=private_error(exc))
            break
        except Exception as exc:
            add(task,'目录扫描/处理失败',note=private_error(exc))
    if not records:write_report(work/'download_report.xlsx',records)
    return records
