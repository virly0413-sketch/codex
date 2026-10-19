import contextlib
import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import blocked_monitor as monitor


class MonitorTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.output=Path(self.temp.name)
        self.config={'channel_id':'UC_test','oauth_token':'token.ru.json'}

    def tearDown(self):self.temp.cleanup()

    def video(self,video_id='abc12345678',region=None,upload='processed',privacy='private'):
        result={'id':video_id,'snippet':{'channelId':'UC_test','title':'剧名'},
                'status':{'uploadStatus':upload,'privacyStatus':privacy},'contentDetails':{}}
        if region is not None:result['contentDetails']['regionRestriction']=region
        return result

    def api(self,videos,missing=None):
        api=MagicMock()
        api.channels.return_value.list.return_value.execute.return_value={'items':[
            {'id':'UC_test','snippet':{'title':'测试频道'},'contentDetails':{'relatedPlaylists':{'uploads':'UU_test'}}}]}
        ids=[v['id'] for v in videos]+list(missing or [])
        pages=[]
        for offset in range(0,max(len(ids),1),50):
            page={'items':[{'contentDetails':{'videoId':i},'snippet':{'title':'列表中的名称'}} for i in ids[offset:offset+50]]}
            if offset+50<len(ids):page['nextPageToken']='page2'
            pages.append(page)
        api.playlistItems.return_value.list.return_value.execute.side_effect=pages
        responses=[]
        for offset in range(0,len(ids),50):
            batch=set(ids[offset:offset+50]);responses.append({'items':[v for v in videos if v['id'] in batch]})
        api.videos.return_value.list.return_value.execute.side_effect=responses
        return api

    def scan(self,api):
        with contextlib.redirect_stdout(io.StringIO()):
            return monitor.scan(self.config,self.output,youtube=api,copyright_only=False)

    def test_private_unlisted_embed_disabled_and_licensed_not_bans(self):
        for privacy in ['private','unlisted','public']:
            video=self.video(privacy=privacy)
            video['status']['embeddable']=False
            video['contentDetails']['licensedContent']=True
            self.assertEqual(monitor.classify(video),[])

    def test_empty_allow_is_worldwide_empty_block_is_not(self):
        self.assertEqual(monitor.classify(self.video(region={'allowed':[]})),[('禁播','全球禁播：允许地区列表为空')])
        self.assertEqual(monitor.classify(self.video(region={'blocked':[]})),[])
        self.assertIn('US',monitor.classify(self.video(region={'blocked':['US']}))[0][1])
        self.assertIn('仅允许',monitor.classify(self.video(region={'allowed':['RU']}))[0][1])

    def test_reject_failure_age_and_deleted_are_distinct(self):
        video=self.video(upload='rejected');video['status']['rejectionReason']='copyright'
        self.assertEqual(monitor.classify(video),[('上传被拒','YouTube拒绝上传：copyright')])
        self.assertEqual(monitor.classify(self.video(upload='failed'))[0][0],'上传失败')
        self.assertEqual(monitor.classify(self.video(upload='deleted'))[0][0],'已删除')
        video=self.video();video['contentDetails']['contentRating']={'ytRating':'ytAgeRestricted'}
        self.assertEqual(monitor.classify(video)[0][0],'年龄限制')

    def test_pagination_missing_detail_and_read_only_calls(self):
        videos=[self.video(f'{n:011d}',region={'blocked':['US']} if n==50 else None) for n in range(51)]
        api=self.api(videos,missing=['missing1234'])
        report,folder=self.scan(api)
        self.assertEqual(report['upload_list_count'],52)
        self.assertEqual(report['details_returned'],51)
        self.assertEqual(report['blocked_video_count'],1)
        self.assertEqual(report['attention'][0]['分类'],'状态待核对')
        self.assertEqual(api.playlistItems.return_value.list.call_count,2)
        self.assertEqual(api.videos.return_value.list.call_count,2)
        self.assertTrue((folder/'blocked_videos.csv').read_bytes().startswith(b'\xef\xbb\xbf'))
        self.assertTrue(all('update' not in str(c) and 'insert' not in str(c) and 'delete' not in str(c) for c in api.mock_calls))

    def test_repeated_scans_and_changed_reason_new_findings(self):
        video=self.video(region={'blocked':['US']})
        first,_=self.scan(self.api([video]));second,_=self.scan(self.api([video]))
        self.assertEqual(len(first['newly_detected']),1)
        self.assertEqual(second['newly_detected'],[])
        video['contentDetails']['regionRestriction']['blocked']=['RU']
        third,_=self.scan(self.api([video]))
        self.assertEqual(len(third['newly_detected']),1)

    def test_scan_failure_preserves_previous_success(self):
        _,folder=self.scan(self.api([self.video(region={'allowed':[]})]))
        prior=(folder/'latest.json').read_bytes()
        old_csv=(folder/'blocked_videos.csv').read_bytes()
        old_excel=(folder/'blocked_videos.xlsx').read_bytes()
        api=self.api([self.video()]);api.videos.return_value.list.return_value.execute.side_effect=TimeoutError()
        with self.assertRaises(RuntimeError):self.scan(api)
        self.assertEqual((folder/'latest.json').read_bytes(),prior)
        self.assertEqual((folder/'blocked_videos.csv').read_bytes(),old_csv)
        self.assertEqual((folder/'blocked_videos.xlsx').read_bytes(),old_excel)

    def test_untrusted_titles_csv_formula_and_html_escape(self):
        video=self.video(region={'allowed':[]});video['snippet']['title']='=HYPERLINK("x")<script>alert(1)</script>'
        _,folder=self.scan(self.api([video]))
        self.assertIn("'=HYPERLINK",(folder/'blocked_videos.csv').read_text(encoding='utf-8-sig'))
        page=(folder/'report.html').read_text()
        self.assertNotIn('<script>',page)
        self.assertIn('&lt;script&gt;',page)

    def test_wrong_owner_aborts(self):
        video=self.video();video['snippet']['channelId']='other_channel'
        with self.assertRaises(ValueError):self.scan(self.api([video]))
        self.assertFalse((self.output/'UC_test'/'latest.json').exists())

    def test_auth_uses_separate_readonly_token(self):
        original=copy.deepcopy(self.config)
        with patch.object(monitor.run,'auth') as auth:
            monitor.monitor_auth(self.config,reauthorize=True)
        self.assertEqual(auth.call_args.args[0]['oauth_token'],'token.ru.monitor.json')
        self.assertEqual(auth.call_args.kwargs['scopes'],['https://www.googleapis.com/auth/youtube.readonly'])
        self.assertEqual(self.config,original)

    def test_excel_two_columns_unique_videos_and_literal_title(self):
        from openpyxl import load_workbook
        video=self.video(region={'allowed':[]},upload='rejected')
        video['snippet']['title']='=HYPERLINK("bad")'
        _,folder=self.scan(self.api([video]))
        book=load_workbook(folder/'blocked_videos.xlsx')
        try:
            sheet=book.active
            self.assertEqual(list(sheet.values),[('视频名称','视频链接'),
                ('=HYPERLINK("bad")','https://www.youtube.com/watch?v=abc12345678')])
            self.assertEqual(sheet['A2'].data_type,'s')
            self.assertEqual(sheet['B2'].hyperlink.target,'https://www.youtube.com/watch?v=abc12345678')
        finally:book.close()

    def test_excel_no_findings_contains_only_headers(self):
        from openpyxl import load_workbook
        _,folder=self.scan(self.api([self.video()]))
        book=load_workbook(folder/'blocked_videos.xlsx')
        try:self.assertEqual(list(book.active.values),[('视频名称','视频链接')])
        finally:book.close()

    def test_command_runs_exactly_one_scan(self):
        with patch('sys.argv',['blocked_monitor.py']),patch.object(monitor.run,'load',return_value=self.config),\
                patch.object(monitor,'scan') as scan:
            self.assertEqual(monitor.main(),0)
        scan.assert_called_once_with(self.config,None)


if __name__=='__main__':unittest.main()
