"""Read-only YouTube restriction monitor. No video updates or uploads."""
import argparse
import csv
import datetime as dt
import html
import io
import json
from pathlib import Path
import sys
from urllib.parse import quote
from zoneinfo import ZoneInfo

import run
from networking import read_request

LIMITATION = ('普通YouTube API不提供完整Content ID版权声明/Studio限制详情；'
              '也不提供黄标、广告适宜性或完整收益状态；需另行登录Studio检查收益。未返回禁播标记不代表无限制。私密、未列出和关闭嵌入不作为禁播。')
FIELDS = ['频道名称', '频道ID', '视频名称', '视频链接', 'Studio链接', '分类', '限制详情',
          '上传状态', '可见性', '检测时间']


def classify(video):
    status = video.get('status', {})
    details = video.get('contentDetails', {})
    findings = []
    region = details.get('regionRestriction', {})
    if 'allowed' in region:
        allowed = region['allowed']
        findings.append(('禁播', '全球禁播：允许地区列表为空' if not allowed
                         else '地区禁播：仅允许以下地区观看：' + ', '.join(allowed)))
    if region.get('blocked'):
        findings.append(('禁播', '地区禁播：禁止观看地区：' + ', '.join(region['blocked'])))
    upload = status.get('uploadStatus')
    if upload == 'rejected':
        findings.append(('上传被拒', 'YouTube拒绝上传：' + status.get('rejectionReason', '原因未返回')))
    elif upload == 'failed':
        findings.append(('上传失败', '处理/上传失败：' + status.get('failureReason', '原因未返回')))
    elif upload == 'deleted':
        findings.append(('已删除', 'API返回deleted；无法据此认定为版权禁播'))
    if details.get('contentRating', {}).get('ytRating') == 'ytAgeRestricted':
        findings.append(('年龄限制', '仅限符合年龄要求的观众；可能影响广告投放，实际收益状态需在Studio核对；不等同于全球禁播'))
    return findings


def read_channel(youtube, channel_id):
    channels = read_request(youtube.channels().list(part='snippet,contentDetails', id=channel_id),
                            '读取频道上传列表').get('items', [])
    if len(channels) != 1 or channels[0]['id'] != channel_id:
        raise ValueError('未能读取目标频道，未生成成功报告。')
    channel = channels[0]
    playlist = channel['contentDetails']['relatedPlaylists']['uploads']
    ids, fallback, seen_pages = [], {}, set()
    page = None
    while True:
        response = read_request(youtube.playlistItems().list(part='snippet,contentDetails',
            playlistId=playlist, maxResults=50, pageToken=page), '读取上传列表')
        for item in response.get('items', []):
            video_id = item.get('contentDetails', {}).get('videoId')
            if video_id:
                ids.append(video_id)
                fallback[video_id] = item.get('snippet', {}).get('title', '名称未返回')
        print('已读取上传列表条目：', len(ids), flush=True)
        page = response.get('nextPageToken')
        if not page:
            break
        if page in seen_pages:
            raise ValueError('API分页令牌重复，扫描未完成；未更新最近成功报告。')
        seen_pages.add(page)
    ids = list(dict.fromkeys(ids))
    videos = {}
    for offset in range(0, len(ids), 50):
        print(f'读取限制状态：{offset + 1}—{min(offset + 50, len(ids))}/{len(ids)}', flush=True)
        response = read_request(youtube.videos().list(part='snippet,status,contentDetails',
            id=','.join(ids[offset:offset + 50])), '读取视频限制状态')
        for video in response.get('items', []):
            if video.get('snippet', {}).get('channelId') != channel_id:
                raise ValueError('API返回其他频道视频，已停止，未更新报告。')
            videos[video['id']] = video
    return channel, ids, videos, fallback


def csv_text(rows):
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=FIELDS)
    writer.writeheader()
    for row in rows:
        # Titles are untrusted text. Prevent formulas when opening the CSV in Excel.
        writer.writerow({k: "'" + str(v) if str(v).lstrip().startswith(('=', '+', '-', '@'))
                         or str(v).startswith(('\t', '\r', '\n')) else v for k, v in row.items()})
    return stream.getvalue()


def atomic_text(path, body, encoding='utf-8'):
    import os
    import tempfile
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding=encoding, newline='') as out:
            out.write(body)
        os.replace(temp, path)
    finally:
        if Path(temp).exists(): Path(temp).unlink()


def export_excel(path, rows, studio=None):
    """Export one row per video, using literal strings and clickable links."""
    import os
    import tempfile
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = '需处理视频'
    sheet.append(['视频名称', '视频链接'])
    seen = set()
    for row in rows:
        url = row['视频链接']
        if url in seen:
            continue
        seen.add(url)
        sheet.append([row['视频名称'], url])
        title_cell, url_cell = sheet[sheet.max_row]
        title_cell.data_type = 's'
        url_cell.data_type = 's'
        url_cell.hyperlink = url
        url_cell.style = 'Hyperlink'
        title_cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
    sheet.column_dimensions['A'].width = 70
    sheet.column_dimensions['B'].width = 65
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    detail = book.create_sheet('限制详情')
    detail.append(['视频名称', '视频链接', '分类', '限制详情'])
    for row in rows:
        detail.append([row.get(k, '') for k in ('视频名称', '视频链接', '分类', '限制详情')])
        for cell in detail[detail.max_row]: cell.data_type = 's'
        detail.cell(detail.max_row, 2).hyperlink = row['视频链接']
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    for col, width in [('A',60),('B',60),('C',28),('D',100)]: detail.column_dimensions[col].width=width
    coverage = book.create_sheet('检测范围')
    coverage.append(['项目', '结果'])
    coverage.append(['普通API', LIMITATION])
    coverage.append(['Studio收益检查', ('已完成，检查'+str(studio['checked'])+'个视频') if studio is not None else '未执行，不能判断是否存在收入限制'])
    coverage.append(['收益说明', '收益关闭、审核中、未知状态均登记为待核对；不等同于确认收入受限。Studio页面变化可能导致无法识别。'])
    coverage.column_dimensions['A'].width=25;coverage.column_dimensions['B'].width=130
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(dir=path.parent, suffix='.xlsx')
    os.close(fd)
    try:
        book.save(temp)
        os.replace(temp, path)
    except PermissionError as exc:
        raise PermissionError('无法保存Excel，请先关闭已打开的blocked_videos.xlsx，并确认报告目录可写。') from exc
    finally:
        book.close()
        if Path(temp).exists(): Path(temp).unlink()


def monitor_auth(config, reauthorize=False):
    config = dict(config)
    original = Path(config['oauth_token'])
    # Separate read-only grant; never overwrite the publishing program's token.
    config['oauth_token'] = str(original.with_name(original.stem + '.monitor.json'))
    return run.auth(config, reauthorize=reauthorize,
                    scopes=['https://www.googleapis.com/auth/youtube.readonly'])


def scan(config, output=None, youtube=None, studio=None):
    if youtube is None:
        youtube = monitor_auth(config)
    channel, ids, videos, fallback = read_channel(youtube, config['channel_id'])
    checked = dt.datetime.now(ZoneInfo('Asia/Shanghai'))
    when = checked.isoformat(timespec='seconds')
    rows = []
    for video_id in ids:
        video = videos.get(video_id)
        findings = classify(video) if video else [('状态待核对', '上传列表中存在，但详情未返回；不能认定禁播或删除')]
        if not findings:
            findings = [('API未返回禁播标记', '不代表无版权限制；必要时在Studio核对')]
        for category, reason in findings:
            safe_id = quote(video_id, safe='')
            rows.append(dict(zip(FIELDS, [channel['snippet']['title'], config['channel_id'],
                video.get('snippet', {}).get('title', fallback[video_id]) if video else fallback[video_id],
                'https://www.youtube.com/watch?v=' + safe_id,
                'https://studio.youtube.com/video/' + safe_id + '/edit', category, reason,
                video.get('status', {}).get('uploadStatus', '') if video else '',
                video.get('status', {}).get('privacyStatus', '') if video else '', when])))
    revenue = []
    if studio is not None:
        for item in studio['rows']:
            row = {key: '' for key in FIELDS}
            row.update({key: item[key] for key in ('视频名称','视频链接','分类','限制详情')})
            row.update({'频道名称':channel['snippet']['title'],'频道ID':config['channel_id'],'检测时间':when})
            revenue.append(row)
    blocked = [r for r in rows if r['分类'] in ('禁播', '上传被拒')]
    attention = [r for r in rows if r['分类'] not in ('禁播', '上传被拒', 'API未返回禁播标记')]
    folder = (Path(output) if output else run.ROOT / 'blocked_reports') / config['channel_id']
    folder.mkdir(parents=True, exist_ok=True)
    previous_path = folder / 'latest.json'
    previous = run.load(previous_path) if previous_path.exists() else {}
    signature = lambda r: (r['视频链接'], r['分类'], r['限制详情'])
    prior = {signature(r) for r in previous.get('blocked', [])}
    new = [r for r in blocked if signature(r) not in prior]
    report = dict(channel_id=config['channel_id'], channel_name=channel['snippet']['title'],
        checked_at=when, upload_list_count=len(ids), details_returned=len(videos),
        blocked_video_count=len({r['视频链接'] for r in blocked}), blocked=blocked,
        attention=attention, revenue=revenue, studio_checked=studio['checked'] if studio is not None else None,
        excel_video_count=len({r['视频链接'] for r in blocked+attention+revenue}),
        newly_detected=new, limitation=LIMITATION,
        scope='仅频道上传列表中API仍返回的条目；不覆盖已从上传列表消失的视频',
        all_videos=rows)
    stamp = checked.strftime('%Y%m%d_%H%M%S_%f')
    body = json.dumps(report, ensure_ascii=False, indent=2)
    # Complete every remote read before replacing the last successful scan.
    export_excel(folder / 'blocked_videos.xlsx', blocked + attention + revenue, studio)
    atomic_text(folder / f'scan_{stamp}.json', body)
    for filename, data in [('blocked_videos.csv', blocked), ('attention.csv', attention),
                           ('newly_detected.csv', new), ('all_videos.csv', rows), ('revenue.csv', revenue)]:
        atomic_text(folder / filename, csv_text(data), 'utf-8-sig')
    rendered = ''.join('<tr>' + ''.join('<td>' + (f'<a href="{html.escape(r[k], quote=True)}" target="_blank" rel="noopener noreferrer">打开</a>'
        if k in ('视频链接', 'Studio链接') else html.escape(str(r[k]))) + '</td>'
        for k in FIELDS) + '</tr>' for r in blocked + attention + revenue)
    page = ('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>YouTube限制检测报告</title>'
        '<style>body{font:15px system-ui;margin:24px}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:8px}th{background:#edf2f7}</style>'
        '<h1>YouTube限制检测报告</h1><p>' + html.escape(channel['snippet']['title'] + ' | ' + when) + '</p>'
        f'<p>Excel共{report["excel_video_count"]}个需处理/待核对视频；Studio收益检查：{report["studio_checked"] if studio is not None else "未执行"}。上传列表：{len(ids)}条；禁播/被拒：{report["blocked_video_count"]}个视频；本次新增或原因变更：{len(new)}条记录。</p>'
        '<p>' + html.escape(LIMITATION) + '</p><p>' + html.escape(report['scope']) + '</p><table><tr>' +
        ''.join('<th>' + html.escape(k) + '</th>' for k in FIELDS) + '</tr>' + rendered + '</table>'
        '<p>没有列出限制记录时，仅表示本次API未返回相关标记。</p></html>')
    atomic_text(folder / 'report.html', page)
    atomic_text(previous_path, body)
    print(f'完成：{len(ids)}条上传记录，{report["blocked_video_count"]}个视频返回禁播/被拒标记。')
    print('报告目录：', folder.resolve())
    print('Excel文件：', (folder / 'blocked_videos.xlsx').resolve())
    print(LIMITATION)
    return report, folder.resolve()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', default=str(run.ROOT / 'config.json'))
    parser.add_argument('--output', type=Path)
    parser.add_argument('--reauth', action='store_true', help='重新选择账号/品牌频道')
    parser.add_argument('--studio', action='store_true', help='登录Studio读取普通视频收益列')
    parser.add_argument('--studio-browser', default='auto', choices=['auto','msedge','chrome','chromium'])
    args = parser.parse_args()
    config = run.load(args.config)
    if args.reauth: monitor_auth(config, reauthorize=True)
    try:
        if args.studio:
            from studio_revenue import scan_studio
            studio = scan_studio(config['channel_id'], run.ROOT/'local_work_studio_monitor'/config['channel_id'], args.studio_browser)
            scan(config, args.output, studio=studio)
        else:
            scan(config, args.output)
    except Exception as exc:
        print(f'扫描未完成，未标记成功：{exc}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print('监测已停止。')
