import csv
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from share_check import schedule_tasks, audit, check
from network_upload import inventory

class ShareCheckTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.schedule=self.root/'schedule.csv';self.share=self.root/'share';self.share.mkdir()
        with self.schedule.open('w',encoding='utf-8-sig',newline='') as f:
            csv.writer(f).writerows([['日期','2026.10.26','2026.10.27','2026.10.28'],['俄语男频','甲剧长名','乙剧长名','丙剧长名'],['英语男频','甲剧长名','','']])
        self.selections=[{'row':2,'language':'俄语'}]
    def add(self,path):
        p=self.share/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(b'video');return p
    def test_realistic_share_matches_language_and_reports_other_audio(self):
        self.add('甲剧长名俄语.mp4');self.add('乙剧长名英语.mp4')
        results=check(self.schedule,self.share,self.selections,self.root/'report.xlsx','2026-10-26',3,progress=lambda s:None)
        self.assertEqual([r['status'] for r in results],['已找到','同剧其他语种或语种未标明','未匹配对应视频'])
        self.assertEqual(len(results[1]['related']),1)
        b=load_workbook(self.root/'report.xlsx');self.addCleanup(b.close)
        self.assertEqual(b['未匹配对应语种'].max_row,3)
        self.assertEqual(b['已找到'].max_row,2)
        self.assertEqual(b['全部排片']['E2'].value,'B2')
    def test_multiple_files_not_misreported_missing(self):
        self.add('甲剧长名/俄语/上.mp4');self.add('甲剧长名/俄语/下.mp4')
        tasks=schedule_tasks(self.schedule,self.selections,'2026-10-26',1)
        r=audit(tasks,inventory(self.share))[0]
        self.assertEqual(r['status'],'多个对应视频');self.assertEqual(len(r['files']),2)
    def test_similar_name_only_suggestion(self):
        self.add('甲剧长名续集俄语.mp4')
        r=audit(schedule_tasks(self.schedule,self.selections,'2026-10-26',1),inventory(self.share))[0]
        self.assertEqual(r['files'],[]);self.assertEqual(r['status'],'疑似同剧，需核对命名')
    def test_schedule_channel_and_calendar_window(self):
        tasks=schedule_tasks(self.schedule,[{'row':3,'language':'英语'}],'2026-10-26',2)
        self.assertEqual(len(tasks),1);self.assertEqual(tasks[0]['source_cell'],'B3')
        with self.assertRaises(ValueError):schedule_tasks(self.schedule,[{'row':2,'language':'英语'}])
        with self.assertRaises(ValueError):schedule_tasks(self.schedule,self.selections,'',3)
    def test_access_failure_preserves_existing_report(self):
        report=self.root/'report.xlsx';report.write_bytes(b'old successful report')
        with patch('share_check.inventory',side_effect=PermissionError('no access')):
            with self.assertRaises(PermissionError):check(self.schedule,self.share,self.selections,report,progress=lambda s:None)
        self.assertEqual(report.read_bytes(),b'old successful report')
    def test_stop_prevents_export(self):
        stop=threading.Event();stop.set();report=self.root/'report.xlsx'
        with self.assertRaises(InterruptedError):check(self.schedule,self.share,self.selections,report,stop=stop,progress=lambda s:None)
        self.assertFalse(report.exists())
    def test_default_previous_schedule_has_84_language_tasks(self):
        from baidu_download import schedule_rows
        source=Path(__file__).parent/'schedules/2026-10-26_to_2026-11-15_7day/schedule_transposed.csv'
        rows=schedule_rows(source);tasks=schedule_tasks(source,rows,'2026-10-26',21)
        self.assertEqual(len(tasks),84)
        self.assertEqual({t['language'] for t in tasks},{'韩语','俄语','西语','菲律宾语'})
        kr=[t for t in tasks if t['language']=='韩语']
        self.assertEqual([t['title'] for t in kr[:3]],['代号枪神','满级大师兄制霸都市','老年逆袭我竟成为山野之王'])
    def test_xlsx_title_is_literal_even_with_formula_characters(self):
        with self.schedule.open('w',encoding='utf-8-sig',newline='') as f:
            csv.writer(f).writerows([['日期','2026.10.26'],['俄语男频','=1+2']])
        check(self.schedule,self.share,self.selections,self.root/'report.xlsx',progress=lambda s:None)
        b=load_workbook(self.root/'report.xlsx');self.addCleanup(b.close)
        self.assertEqual(b['未匹配对应语种']['D2'].value,'=1+2')
        self.assertEqual(b['未匹配对应语种']['D2'].data_type,'s')
