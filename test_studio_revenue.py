import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from studio_revenue import revenue_status
import test_blocked_monitor as fixtures

class RevenueTests(unittest.TestCase):
    def test_explicit_limits(self):
        for label in ('Limited','Limited or no ads','受限','广告受限','Ограничено'):
            self.assertEqual(revenue_status(label)[0],'收入限制')
        self.assertEqual(revenue_status('Ineligible')[0],'无法创收')
    def test_off_is_not_confirmed_limit(self):
        self.assertEqual(revenue_status('Off')[0],'收益关闭（需核对）')
        self.assertEqual(revenue_status('Checking')[0],'收益审核中')
        self.assertEqual(revenue_status('Sharing')[0],'收益分成')
    def test_unknown_not_healthy(self):
        for label in ('','unrecognized icon','Not available'):
            self.assertEqual(revenue_status(label)[0],'收益待核对')
        self.assertIsNone(revenue_status('On')[0])
    def test_combined_excel_details_and_dedup(self):
        import blocked_monitor as monitor
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        video=helper.video(region={'blocked':['US']});video['snippet']['title']='=unsafe'
        studio={'checked':2,'rows':[
            {'视频名称':'=unsafe','视频链接':'https://www.youtube.com/watch?v=abc12345678','分类':'收入限制','限制详情':'Limited'},
            {'视频名称':'second','视频链接':'https://www.youtube.com/watch?v=other123456','分类':'收益待核对','限制详情':'未显示'}]}
        with contextlib.redirect_stdout(io.StringIO()):
            report,folder=monitor.scan(helper.config,helper.output,youtube=helper.api([video]),studio=studio)
        self.assertEqual(report['excel_video_count'],2)
        self.assertEqual(report['studio_checked'],2)
        book=load_workbook(folder/'blocked_videos.xlsx')
        self.addCleanup(book.close)
        self.assertEqual(book.active.max_row,3)
        self.assertEqual(book.active['A2'].data_type,'s')
        self.assertEqual(book['限制详情'].max_row,4)
        self.assertIn('已完成',book['检测范围']['B3'].value)
        self.assertIn('Limited',(folder/'revenue.csv').read_text(encoding='utf-8-sig'))
    def test_api_only_scope_never_implies_revenue_ok(self):
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        report,folder=helper.scan(helper.api([helper.video()]))
        self.assertIsNone(report['studio_checked'])
        book=load_workbook(folder/'blocked_videos.xlsx');self.addCleanup(book.close)
        self.assertIn('未执行',book['检测范围']['B3'].value)

if __name__=='__main__':unittest.main()
