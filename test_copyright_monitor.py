import contextlib
import io
import unittest
from openpyxl import load_workbook
import blocked_monitor as monitor
import test_blocked_monitor as fixtures

class CopyrightTests(unittest.TestCase):
    def test_only_explicit_copyright_is_registered(self):
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        copyright=helper.video('copyright01',upload='rejected')
        copyright['status']['rejectionReason']='copyright'
        legal=helper.video('legal000001',upload='rejected');legal['status']['rejectionReason']='legal'
        age=helper.video('age00000001');age['contentDetails']['contentRating']={'ytRating':'ytAgeRestricted'}
        region=helper.video('region00001',region={'blocked':['US']})
        licensed=helper.video('license0001');licensed['contentDetails']['licensedContent']=True
        failed=helper.video('failed00001',upload='failed')
        videos=[copyright,legal,age,region,licensed,failed]
        with contextlib.redirect_stdout(io.StringIO()):
            report,folder=monitor.scan(helper.config,helper.output,youtube=helper.api(videos),copyright_only=True)
        self.assertEqual(report['excel_video_count'],1)
        self.assertEqual(len(report['registered']),1)
        self.assertIn('copyright01',report['registered'][0]['视频链接'])
        book=load_workbook(folder/'blocked_videos.xlsx');self.addCleanup(book.close)
        self.assertEqual(book.active.title,'版权问题')
        self.assertEqual(book.active.max_row,2)
        self.assertEqual(book['限制详情'].max_row,2)
        self.assertIn('copyright',book['限制详情']['D2'].value)
        self.assertIn('copyright01',(folder/'copyright_videos.csv').read_text(encoding='utf-8-sig'))
    def test_region_block_is_not_proof_of_copyright(self):
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        with contextlib.redirect_stdout(io.StringIO()):
            report,folder=monitor.scan(helper.config,helper.output,youtube=helper.api([helper.video(region={'allowed':[]})]),copyright_only=True)
        self.assertEqual(report['registered'],[])
        self.assertIn('空名单不代表没有版权问题',report['limitation'])
