import contextlib
import io
import unittest
from openpyxl import load_workbook
import blocked_monitor as monitor
import test_blocked_monitor as fixtures

class RegistrationTests(unittest.TestCase):
    def test_bans_and_evidenced_revenue_risks_both_exported(self):
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        blocked=helper.video('blocked0001',region={'blocked':['US']})
        rejected=helper.video('reject00001',upload='rejected');rejected['status']['rejectionReason']='copyright'
        age=helper.video('agerisk0001');age['contentDetails']['contentRating']={'ytRating':'ytAgeRestricted'}
        failed=helper.video('failed00001',upload='failed')
        licensed=helper.video('licensed001');licensed['contentDetails']['licensedContent']=True
        with contextlib.redirect_stdout(io.StringIO()):
            report,folder=monitor.scan(helper.config,helper.output,youtube=helper.api([blocked,rejected,age,failed,licensed]))
        self.assertEqual({r['分类'] for r in report['registered']},{'禁播','上传被拒','年龄限制'})
        self.assertEqual(report['excel_video_count'],3)
        self.assertIsNone(report['studio_checked'])
        book=load_workbook(folder/'blocked_videos.xlsx');self.addCleanup(book.close)
        self.assertEqual(book.active.max_row,4)
        self.assertEqual(book.active.max_column,2)
        self.assertEqual(book['限制详情'].max_row,4)
        reasons=' '.join(str(r[3]) for r in list(book['限制详情'].values)[1:])
        self.assertIn('实际收益状态需在Studio核对',reasons)
    def test_combined_restrictions_deduplicate_video(self):
        helper=fixtures.MonitorTests();helper.setUp();self.addCleanup(helper.tearDown)
        video=helper.video(region={'allowed':[]})
        video['contentDetails']['contentRating']={'ytRating':'ytAgeRestricted'}
        with contextlib.redirect_stdout(io.StringIO()):
            report,folder=monitor.scan(helper.config,helper.output,youtube=helper.api([video]))
        self.assertEqual(report['excel_video_count'],1)
        book=load_workbook(folder/'blocked_videos.xlsx');self.addCleanup(book.close)
        self.assertEqual(book.active.max_row,2)
        self.assertEqual(book['限制详情'].max_row,3)
